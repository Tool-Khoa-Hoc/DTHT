"""
word2md.py — Word (.docx) → Markdown Preprocessor
===================================================
Duyệt toàn bộ sub-folder trong input/, đọc file .docx, áp dụng các quy tắc
chuẩn hóa và xuất Markdown sạch vào output/<tên-folder>/<tên-file>.md

Chế độ thủ công (mặc định):
  1. Loại bỏ nhiễu (số trang, header/footer)
  2. Chuyển bảng docx → Markdown table
  3. Nhận diện PHẦN I/II... → ## Header
  4. Sửa lỗi dính chữ (Câu1 → Câu 1, v.v.)
  5. Giữ nguyên [img:...] và [math:...]
  6. Xóa khoảng trắng thừa

Chế độ AI (--ai flag):
  Dùng lib/raw_extractor.py để trích xuất cấu trúc + ảnh,
  sau đó gọi Gemini API để làm sạch nội dung.

Usage:
  python word2md.py                      # Xử lý toàn bộ input/ (thủ công)
  python word2md.py --ai --key <api_key> # Dùng Gemini AI
  python word2md.py --file path          # Xử lý 1 file cụ thể
  python word2md.py --preview            # In ra console (không ghi file)
"""

from __future__ import annotations

import argparse
import re
import sys
import unicodedata
from pathlib import Path

try:
    from docx import Document
    from docx.oxml.ns import qn
except ImportError:
    print("[ERROR] Thiếu thư viện 'python-docx'. Chạy: pip install python-docx")
    sys.exit(1)

# ──────────────────────────────────────────────
# Constants
# ──────────────────────────────────────────────
INPUT_DIR  = Path(__file__).parent / "input"
OUTPUT_DIR = Path(__file__).parent / "output"

NOISE_PATTERNS = [
    re.compile(r"^\s*\d+\s*$"),
    re.compile(r"(?i)^\s*trang\s*\d+\s*$"),
    re.compile(r"(?i)^\s*page\s*\d+\s*$"),
    re.compile(r"(?i)^\s*-\s*\d+\s*-\s*$"),
    re.compile(r"(?i)^\s*\(trang\s*\d+\)\s*$"),
]

SECTION_RE = re.compile(
    r"^(phần\s+[IVXivx]+\s*[:\.\-]?.*|PHẦN\s+[IVXivx]+\s*[:\.\-]?.*)$",
    re.IGNORECASE,
)

TYPO_FIXES = [
    (re.compile(r"(Câu|câu)(\d+)"), r"\1 \2"),
]

PROTECTED_TAG_RE = re.compile(r"(\[img:[^\]]+\]|\[math:[^\]]+\])")

# ──────────────────────────────────────────────
# Helpers (chế độ thủ công)
# ──────────────────────────────────────────────

def is_noise(text: str) -> bool:
    t = text.strip()
    if not t:
        return False
    for pat in NOISE_PATTERNS:
        if pat.search(t):
            return True
    return False


def fix_typos(text: str) -> str:
    for pat, repl in TYPO_FIXES:
        text = pat.sub(repl, text)
    return text


def normalize_spaces(text: str) -> str:
    parts = PROTECTED_TAG_RE.split(text)
    result = []
    for i, part in enumerate(parts):
        if i % 2 == 1:
            result.append(part)
        else:
            part = re.sub(r"[ \t]+", " ", part)
            result.append(part)
    return "".join(result).strip()


def convert_section_header(text: str) -> str | None:
    t = text.strip()
    if SECTION_RE.match(t):
        return f"## {t.upper()}"
    return None


def table_to_markdown(table) -> str:
    rows = []
    for row in table.rows:
        cells = []
        for cell in row.cells:
            cell_text = " ".join(p.text for p in cell.paragraphs).strip()
            cell_text = cell_text.replace("|", "\\|")
            cells.append(cell_text)
        rows.append(cells)

    if not rows:
        return ""

    seen = []
    unique_rows = []
    for row in rows:
        key = tuple(row)
        if key not in seen:
            seen.append(key)
            unique_rows.append(row)
    rows = unique_rows

    max_cols = max(len(r) for r in rows)
    rows = [r + [""] * (max_cols - len(r)) for r in rows]

    lines = []
    header = "| " + " | ".join(rows[0]) + " |"
    sep    = "| " + " | ".join(["---"] * max_cols) + " |"
    lines.append(header)
    lines.append(sep)
    for row in rows[1:]:
        lines.append("| " + " | ".join(row) + " |")
    return "\n".join(lines)


def process_document(docx_path: Path) -> str:
    doc = Document(str(docx_path))
    output_lines: list[str] = []
    prev_blank = False

    body = doc.element.body
    para_map  = {p._element: p for p in doc.paragraphs}
    table_map = {t._element: t for t in doc.tables}

    for child in body:
        tag = child.tag
        if tag == qn("w:p"):
            para = para_map.get(child)
            if para is None:
                continue
            raw = para.text

            if not raw.strip():
                if not prev_blank and output_lines:
                    output_lines.append("")
                    prev_blank = True
                continue

            if is_noise(raw):
                continue

            header = convert_section_header(raw)
            if header:
                if output_lines and output_lines[-1] != "":
                    output_lines.append("")
                output_lines.append(header)
                output_lines.append("")
                prev_blank = True
                continue

            processed = fix_typos(raw)
            processed = normalize_spaces(processed)

            if processed:
                output_lines.append(processed)
                prev_blank = False

        elif tag == qn("w:tbl"):
            table = table_map.get(child)
            if table is None:
                continue
            md_table = table_to_markdown(table)
            if md_table:
                if output_lines and output_lines[-1] != "":
                    output_lines.append("")
                output_lines.append(md_table)
                output_lines.append("")
                prev_blank = True

    while output_lines and output_lines[-1] == "":
        output_lines.pop()

    return "\n".join(output_lines)


def process_folder(input_folder: Path, output_base: Path, preview: bool = False):
    docx_files = list(input_folder.glob("*.docx"))
    if not docx_files:
        print(f"  [skip] Không tìm thấy .docx trong: {input_folder.name}")
        return

    for docx_path in docx_files:
        print(f"  [→] Xử lý: {docx_path.name} ...", end=" ")
        try:
            md_content = process_document(docx_path)
            if preview:
                print()
                print("─" * 60)
                print(md_content[:2000])
                if len(md_content) > 2000:
                    print(f"\n... ({len(md_content) - 2000} ký tự còn lại)")
                print("─" * 60)
            else:
                out_dir = output_base / input_folder.name
                out_dir.mkdir(parents=True, exist_ok=True)
                out_file = out_dir / (docx_path.stem + ".md")
                out_file.write_text(md_content, encoding="utf-8")
                print(f"[OK] → {out_file.relative_to(output_base.parent)}")
        except Exception as e:
            print(f"[LỖI] {e}")


# ──────────────────────────────────────────────
# AI mode integration
# ──────────────────────────────────────────────

def _process_with_ai(
    docx_path: Path,
    out_dir: Path | None,
    api_key: str | None,
    preview: bool,
):
    """Xử lý 1 file .docx dùng raw_extractor + Gemini AI."""
    try:
        sys.path.insert(0, str(Path(__file__).parent))
        from lib.raw_extractor import extract as re_extract
    except ImportError as e:
        print(f"[LỖI] Không import được lib/raw_extractor: {e}")
        return

    print(f"  [AI→] Xử lý: {docx_path.name} ...", end=" ")
    try:
        assets_dir = (out_dir / docx_path.parent.name / "assets") if out_dir else None
        md, image_map = re_extract(
            docx_path=docx_path,
            assets_dir=assets_dir,
            use_ai=True,
            api_key=api_key,
        )
        if preview:
            print()
            print("─" * 60)
            print(md[:2000])
            if len(md) > 2000:
                print(f"\n... ({len(md) - 2000} ký tự còn lại)")
            print("─" * 60)
        else:
            folder = out_dir / docx_path.parent.name if out_dir else docx_path.parent
            folder.mkdir(parents=True, exist_ok=True)
            out_file = folder / (docx_path.stem + ".md")
            out_file.write_text(md, encoding="utf-8")
            imgs = len(image_map)
            print(f"[OK] → {out_file.name}" + (f" ({imgs} ảnh)" if imgs else ""))
    except Exception as e:
        print(f"[LỖI] {e}")


# ──────────────────────────────────────────────
# Main
# ──────────────────────────────────────────────

def main():
    parser = argparse.ArgumentParser(
        description="Chuyển đổi file .docx đề thi sang Markdown chuẩn hóa"
    )
    parser.add_argument(
        "--file", "-f",
        type=Path,
        help="Xử lý 1 file .docx cụ thể thay vì toàn bộ input/"
    )
    parser.add_argument(
        "--preview", "-p",
        action="store_true",
        help="In kết quả ra console (không ghi file)"
    )
    parser.add_argument(
        "--out", "-o",
        type=Path,
        default=OUTPUT_DIR,
        help=f"Thư mục output (mặc định: {OUTPUT_DIR})"
    )
    parser.add_argument(
        "--ai",
        action="store_true",
        help="Dùng Gemini AI (lib/raw_extractor) thay vì Regex thủ công"
    )
    parser.add_argument(
        "--key", "-k",
        default=None,
        help="Gemini API key (hoặc đặt env GEMINI_API_KEY)"
    )
    args = parser.parse_args()

    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")

    if args.file:
        docx_path: Path = args.file.resolve()
        if not docx_path.exists():
            print(f"[ERROR] Không tìm thấy file: {docx_path}")
            sys.exit(1)
        print(f"[word2md] Xử lý file: {docx_path.name}")
        if args.ai:
            _process_with_ai(docx_path, args.out, args.key, args.preview)
        else:
            try:
                md = process_document(docx_path)
            except Exception as e:
                print(f"[LỖI] {e}")
                sys.exit(1)
            if args.preview:
                print(md)
            else:
                args.out.mkdir(parents=True, exist_ok=True)
                out_file = args.out / (docx_path.stem + ".md")
                out_file.write_text(md, encoding="utf-8")
                print(f"[OK] Đã ghi: {out_file}")
        return

    # Xử lý toàn bộ input/
    if not INPUT_DIR.exists():
        print(f"[ERROR] Không tìm thấy thư mục: {INPUT_DIR}")
        sys.exit(1)

    sub_folders = [f for f in INPUT_DIR.iterdir() if f.is_dir()]
    if not sub_folders:
        print("[WARN] Không có sub-folder nào trong input/")
        sys.exit(0)

    mode_label = "AI (Gemini)" if args.ai else "Thủ công (Regex)"
    print(f"[word2md] Tìm thấy {len(sub_folders)} thư mục | Chế độ: {mode_label}\n")

    for folder in sorted(sub_folders):
        print(f"📁 {folder.name}")
        if args.ai:
            for docx_path in sorted(folder.glob("*.docx")):
                _process_with_ai(docx_path, args.out, args.key, args.preview)
        else:
            process_folder(folder, args.out, preview=args.preview)
        print()

    if not args.preview:
        print(f"✅ Hoàn tất! Kết quả trong: {args.out}")


if __name__ == "__main__":
    main()
